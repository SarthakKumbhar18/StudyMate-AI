import os

import pymupdf
from docx import Document
from pptx import Presentation


def extract_pdf_text(file):

    document = pymupdf.open(
    stream=file.read(),
    filetype="pdf"
)

    text = ""

    for page in document:
        text += page.get_text()

    document.close()

    return text


def extract_docx_text(file):

    document = Document(file)

    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    return text


def extract_pptx_text(file):

    presentation = Presentation(file)

    text = ""

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text


def extract_txt_text(file):

    return file.read().decode("utf-8")


def extract_text(file):

    filename = file.name.lower()

    extension = os.path.splitext(filename)[1]

    if extension == ".pdf":

        return extract_pdf_text(file)

    elif extension == ".docx":

        return extract_docx_text(file)

    elif extension == ".pptx":

        return extract_pptx_text(file)

    elif extension == ".txt":

        return extract_txt_text(file)

    else:

        return None